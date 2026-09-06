"""
Text extraction for uploaded protected documents.

Kept out of the endpoint module on purpose: this is pure parsing with no auth,
database or FastAPI dependency, and the endpoint module cannot be imported
without the whole auth stack. Here it is testable on its own, which matters
because CI deliberately does not install the heavy ML requirements.

Parser libraries are imported lazily so the app still starts, and still accepts
pasted documents, when an optional format's package is missing.
"""
import io
import re

# Formats read as plain text. CSV is included because exported client and
# customer tables are one of the likelier things an admin has to hand.
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".log", ".json"}
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# Upload ceiling. A protected document is admin-supplied, but it is still
# untrusted input handed to a third-party parser, so it gets a hard limit
# rather than whatever the client claims.
MAX_UPLOAD_BYTES = 10 * 1024 * 1024

# Below this, treat extraction as having failed rather than storing a document
# the shield silently cannot protect.
MIN_EXTRACTED_CHARS = 20


class ExtractionError(Exception):
    """Extraction failed in a way the admin can act on. `status` is the HTTP code."""

    def __init__(self, status: int, detail: str):
        super().__init__(detail)
        self.status = status
        self.detail = detail


def _pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:                                        # pragma: no cover
        raise ExtractionError(503, "PDF support requires the pypdf package. Paste the text instead.")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError(422, "This PDF is password-protected. Decrypt it before uploading.")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except ExtractionError:
        raise
    except Exception:
        # The parser's own message leaks file internals and rarely tells the
        # admin anything they can act on.
        raise ExtractionError(422, "Could not read this PDF. It may be corrupt or image-only.")


def _docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError:                                        # pragma: no cover
        raise ExtractionError(503, "Word support requires the python-docx package. Paste the text instead.")
    try:
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        # Contract terms and client details usually live in tables rather than
        # paragraphs; skipping them would index half the document.
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except ExtractionError:
        raise
    except Exception:
        raise ExtractionError(
            422,
            "Could not read this Word file. It may be corrupt, or saved as the "
            "older .doc format — re-save it as .docx.",
        )


def extract(filename: str, content_type: str, data: bytes) -> str:
    """
    Return the text of an uploaded document.

    Raises ExtractionError with an HTTP status and a message written for the
    admin who picked the file, not for a log.
    """
    if not data:
        raise ExtractionError(422, "File is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ExtractionError(
            413,
            f"File is {len(data) / 1024 / 1024:.1f}MB; the limit is "
            f"{MAX_UPLOAD_BYTES // 1024 // 1024}MB.",
        )

    suffix = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""

    if suffix == ".pdf" or content_type == "application/pdf":
        text = _pdf(data)
    elif suffix == ".docx" or content_type == DOCX_MIME:
        text = _docx(data)
    elif suffix == ".doc":
        # .doc is a different binary format, not an older .docx — python-docx
        # cannot read it, and the generic error would send the admin nowhere.
        raise ExtractionError(
            415,
            "The legacy .doc format is not supported. Open it in Word and save "
            "as .docx, or paste the text.",
        )
    elif suffix in TEXT_SUFFIXES or content_type.startswith("text/"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ExtractionError(
            415,
            f"Unsupported file type '{suffix or content_type or 'unknown'}'. Upload a "
            f"PDF, a Word .docx, or a text file ({', '.join(sorted(TEXT_SUFFIXES))}).",
        )

    # Collapse the ragged whitespace PDF extraction produces. Entity spans are
    # character offsets into this text, so it is stored exactly as indexed.
    text = re.sub(r"[ \t]+", " ", text).strip()

    # A scanned PDF parses fine and yields nothing.
    if len(text) < MIN_EXTRACTED_CHARS:
        raise ExtractionError(
            422,
            "Almost no text could be extracted. If this is a scanned or "
            "image-only PDF, it needs OCR before it can be protected.",
        )
    return text
